"""US-004 / R4: import XLSX odporny na uszkodzone i wrogie pliki; limit wierszy poza demo."""

import io
import zipfile
from unittest.mock import patch

import openpyxl
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings

from .imports import ImportError_, parse_locations_file
from .models import (
    HandlingUnit,
    IssueDocument,
    Location,
    Profile,
    Role,
    StockMovement,
)

IMPORT_URLS = (
    "/lokalizacje/import/",
    "/stany/import/",
    "/przyjecie/import/",
    "/wydania/import/",
)
XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _zip_bomb(uncompressed_bytes):
    """Poprawny ZIP: mały po kompresji, duży po rozpakowaniu (same zera)."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("xl/worksheets/sheet1.xml", b"\0" * uncompressed_bytes)
    return buf.getvalue()


class _ViewBase(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("op", password="OpPass123!")
        Profile.objects.filter(user=self.user).update(role=Role.ADMIN)
        self.client.login(username="op", password="OpPass123!")

    def _db_snapshot(self):
        return (
            Location.objects.count(),
            HandlingUnit.objects.count(),
            StockMovement.objects.count(),
            IssueDocument.objects.count(),
        )

    def _post(self, url, name, data):
        upload = SimpleUploadedFile(name, data, content_type=XLSX_CT)
        return self.client.post(url, {"file": upload}, follow=True)


class BrokenXlsxViewTests(_ViewBase):
    def test_non_zip_xlsx_rejected_on_all_importers(self):
        """Kryteria 1 i 3: CSV z rozszerzeniem .xlsx (brak sygnatury PK)."""
        before = self._db_snapshot()
        for url in IMPORT_URLS:
            with self.subTest(url=url):
                resp = self._post(url, "dane.xlsx", b"kod;typ\nA-01;ZAPAS\n")
                self.assertEqual(resp.status_code, 200)
                self.assertContains(resp, "Plik nie jest poprawnym XLSX")
        self.assertEqual(self._db_snapshot(), before)

    def test_truncated_zip_rejected_on_all_importers(self):
        """Kryterium 6: sygnatura PK, ale archiwum ucięte (BadZipFile)."""
        data = _xlsx([["kod"], ["A-01"]])[:200]
        before = self._db_snapshot()
        for url in IMPORT_URLS:
            with self.subTest(url=url):
                resp = self._post(url, "dane.xlsx", data)
                self.assertEqual(resp.status_code, 200)
                self.assertContains(resp, "Plik nie jest poprawnym XLSX")
        self.assertEqual(self._db_snapshot(), before)

    @override_settings(IMPORT_MAX_UNCOMPRESSED_BYTES=1024 * 1024)
    def test_zip_bomb_rejected_before_load_workbook(self):
        """Kryterium 2 (limit zmniejszony w teście, żeby nie alokować 50 MB)."""
        data = _zip_bomb(2 * 1024 * 1024)
        self.assertLess(len(data), 50 * 1024)
        with patch("openpyxl.load_workbook") as load:
            for url in IMPORT_URLS:
                with self.subTest(url=url):
                    resp = self._post(url, "bomba.xlsx", data)
                    self.assertContains(resp, "po rozpakowaniu przekracza 1 MB")
            load.assert_not_called()

    def test_valid_xlsx_still_imports(self):
        resp = self._post(
            "/lokalizacje/import/",
            "lok.xlsx",
            _xlsx([["kod", "typ"], ["A-01", "ZAPAS"]]),
        )
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(Location.objects.filter(code="A-01").exists())


class LibraryErrorTests(SimpleTestCase):
    """Kryterium 6: wyjątki biblioteki → ImportError_, nigdy 500."""

    def test_zip_without_workbook_parts(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("readme.txt", "nie arkusz")
        with self.assertRaisesMessage(ImportError_, "Plik nie jest poprawnym XLSX"):
            parse_locations_file("x.xlsx", buf.getvalue())

    def test_library_exceptions_are_wrapped(self):
        from openpyxl.utils.exceptions import InvalidFileException

        data = _xlsx([["kod"], ["A-01"]])
        for exc in (
            KeyError("xl/workbook.xml"),
            InvalidFileException("x"),
            zipfile.BadZipFile("x"),
        ):
            with self.subTest(exc=type(exc).__name__):
                with patch("openpyxl.load_workbook", side_effect=exc):
                    with self.assertRaisesMessage(
                        ImportError_, "Plik nie jest poprawnym XLSX"
                    ):
                        parse_locations_file("x.xlsx", data)

    def test_default_limits(self):
        from django.conf import settings

        self.assertEqual(settings.IMPORT_MAX_ROWS, 10_000)
        self.assertEqual(settings.IMPORT_MAX_UNCOMPRESSED_BYTES, 50 * 1024 * 1024)


@override_settings(DEMO_MODE=False, IMPORT_MAX_ROWS=20)
class RowLimitTests(_ViewBase):
    """Kryterium 4: limit wierszy także poza demo (wartość zmniejszona w teście)."""

    def test_row_limit_enforced_outside_demo(self):
        rows = [["kod"]] + [[f"A-{i:03d}"] for i in range(21)]
        resp = self._post("/lokalizacje/import/", "lok.xlsx", _xlsx(rows))
        self.assertContains(resp, "najwyżej 20 wierszy")
        self.assertEqual(Location.objects.count(), 0)

    def test_row_limit_applies_to_csv(self):
        data = ("kod\n" + "".join(f"B-{i:03d}\n" for i in range(21))).encode()
        with self.assertRaisesMessage(ImportError_, "najwyżej 20 wierszy"):
            parse_locations_file("lok.csv", data)

    def test_boundary_accepted(self):
        rows = [["kod"]] + [[f"A-{i:03d}"] for i in range(20)]
        self._post("/lokalizacje/import/", "lok.xlsx", _xlsx(rows))
        self.assertEqual(Location.objects.count(), 20)

    @override_settings(DEMO_MODE=True, DEMO_IMPORT_MAX_ROWS=5)
    def test_demo_limit_unchanged(self):
        data = ("kod\n" + "".join(f"C-{i}\n" for i in range(6))).encode()
        with self.assertRaisesMessage(
            ImportError_, "W wersji demo import obejmuje najwyżej 5"
        ):
            parse_locations_file("lok.csv", data)


class DefusedXmlTests(SimpleTestCase):
    """Kryterium 5: defusedxml w zależnościach produkcyjnych i wykrywany przez openpyxl."""

    def test_defusedxml_available(self):
        from openpyxl.xml import DEFUSEDXML

        self.assertTrue(DEFUSEDXML)

    def test_defusedxml_pinned_in_requirements(self):
        from pathlib import Path

        req = (Path(__file__).resolve().parents[1] / "requirements.txt").read_text()
        self.assertRegex(req, r"(?m)^defusedxml==\d")
